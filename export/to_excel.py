'''
导出到 Excel - 预算表、材料清单、施工计划
'''
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter


def _styles():
    '''通用样式'''
    return {
        'header_font': Font(name='微软雅黑', bold=True, size=12, color='FFFFFF'),
        'bold_font': Font(name='微软雅黑', bold=True, size=10),
        'normal_font': Font(name='微软雅黑', size=10),
        'header_fill': PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid'),
        'subtotal_fill': PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid'),
        'total_fill': PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid'),
        'thin_border': Border(
            left=Side(style='thin', color='B0B0B0'),
            right=Side(style='thin', color='B0B0B0'),
            top=Side(style='thin', color='B0B0B0'),
            bottom=Side(style='thin', color='B0B0B0'),
        ),
        'center': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'left': Alignment(horizontal='left', vertical='center', wrap_text=True),
        'money_fmt': '#,##0.00',
    }


def export_budget_to_excel(pricing_items: list[dict], output_path: str):
    '''导出预算报价到 Excel'''
    s = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = '预算报价'
    
    widths = [5, 18, 22, 10, 10, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    
    # 标题
    ws.merge_cells('A1:G1')
    c = ws.cell(row=1, column=1, value='室内装修预算报价表')
    c.font = Font(name='微软雅黑', bold=True, size=16, color='1F4E79')
    c.alignment = s['center']
    
    # 表头
    headers = ['序号', '工程类别', '项目名称', '数量', '单位', '单价(元)', '合价(元)']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = s['header_font']; c.fill = s['header_fill']
        c.alignment = s['center']; c.border = s['thin_border']
    
    row = 4
    seq = 0
    current_cat = ''
    subtotals = {}
    
    for item in pricing_items:
        seq += 1
        cat = item.get('category', '')
        if cat != current_cat:
            if current_cat and subtotals.get(current_cat):
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
                c = ws.cell(row=row, column=1,
                            value=f'{current_cat} 小计')
                c.font = s['bold_font']; c.fill = s['subtotal_fill']; c.alignment = s['center']
                c = ws.cell(row=row, column=7, value=subtotals[current_cat])
                c.font = s['bold_font']; c.fill = s['subtotal_fill']
                c.number_format = s['money_fmt']
                for col in range(1, 8):
                    ws.cell(row=row, column=col).border = s['thin_border']
                row += 1
            current_cat = cat
            subtotals[cat] = 0
        
        subtotals[cat] += item.get('total', 0)
        
        values = [seq, cat, item.get('item', ''), item.get('qty', ''),
                  item.get('unit', ''), item.get('price', 0), item.get('total', 0)]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = s['normal_font']; c.border = s['thin_border']
            c.alignment = s['center'] if col in (1, 4, 5) else s['left']
            if col in (6, 7):
                c.number_format = s['money_fmt']
        
        row += 1
    
    # 总计
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    c = ws.cell(row=row, column=1, value='总计')
    c.font = Font(name='微软雅黑', bold=True, size=12); c.fill = s['total_fill']; c.alignment = s['center']
    grand_total = sum(subtotals.values())
    c = ws.cell(row=row, column=7, value=round(grand_total, 2))
    c.font = Font(name='微软雅黑', bold=True, size=12); c.fill = s['total_fill']
    c.number_format = s['money_fmt']
    for col in range(1, 8):
        ws.cell(row=row, column=col).border = s['thin_border']
    
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    wb.save(output_path)
    return output_path


def export_material_list_to_excel(material_data: dict, output_path: str):
    '''导出材料清单到 Excel'''
    s = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = '材料清单'
    
    widths = [5, 12, 18, 14, 8, 10, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells('A1:G1')
    c = ws.cell(row=1, column=1, value='材料清单（含损耗）')
    c.font = Font(name='微软雅黑', bold=True, size=16, color='1F4E79')
    c.alignment = s['center']

    headers = ['序号', '类别', '材料名称', '规格', '数量', '单位', '含损耗']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = s['header_font']; c.fill = s['header_fill']
        c.alignment = s['center']; c.border = s['thin_border']

    summary = material_data.get('summary', [])
    row = 4
    for i, item in enumerate(summary, 1):
        values = [i, item.get('category', ''), item.get('material', ''),
                  item.get('spec', ''), item.get('quantity', ''), item.get('unit', ''),
                  item.get('including_waste', '')]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = s['normal_font']; c.border = s['thin_border']
            c.alignment = s['center'] if col in (1, 4, 5, 6) else s['left']
        row += 1

    wb.save(output_path)
    return output_path


def export_schedule_to_excel(schedule: list[dict], output_path: str):
    '''导出施工计划到 Excel'''
    s = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = '施工计划'

    widths = [5, 16, 14, 14, 8, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells('A1:F1')
    c = ws.cell(row=1, column=1, value='施工进度计划表')
    c.font = Font(name='微软雅黑', bold=True, size=16, color='1F4E79')
    c.alignment = s['center']

    headers = ['序号', '施工阶段', '开始日期', '结束日期', '天数', '工作内容']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = s['header_font']; c.fill = s['header_fill']
        c.alignment = s['center']; c.border = s['thin_border']

    row = 4
    for i, phase in enumerate(schedule, 1):
        values = [i, phase.get('phase', ''), phase.get('start', ''),
                  phase.get('end', ''), phase.get('days', ''), phase.get('description', '')]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = s['normal_font']; c.border = s['thin_border']
            c.alignment = s['center'] if col in (1, 3, 4, 5) else s['left']
        row += 1

    wb.save(output_path)
    return output_path
