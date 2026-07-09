"""
问卷解析器 - Excel 问卷 -> 结构化设计条件
"""
import json, os
from collections import OrderedDict
import openpyxl
from core.data_model import DesignConditions, ProjectInfo, FamilyInfo, StylePreference, BudgetInfo


def parse_survey(filepath: str) -> DesignConditions:
    """解析客户填写的设计需求问卷 Excel，返回 DesignConditions"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    conditions = DesignConditions()

    # ---- Sheet 1: 项目基本信息 ----
    ws1 = wb["项目基本信息"]
    key_map_1 = {
        "项目名称": "name", "项目地址": "address", "房屋类型": "house_type",
        "建筑面积": "area_m2", "户型图": "floor_plan_path", "设计类型": "design_type",
        "完成时间": "expected_completion", "设计师": "designer",
    }
    project = ProjectInfo()
    for r in range(6, 14):
        label = ws1.cell(row=r, column=2).value
        value = ws1.cell(row=r, column=4).value
        if not label:
            continue
        val = value if value and not str(value).startswith("（") else ""
        matched = False
        for cn, en in key_map_1.items():
            if cn in str(label):
                setattr(project, en, val)
                matched = True
                break
        if not matched:
            setattr(project, str(label), val)
    conditions.project = project

    # ---- Sheet 2: 家庭成员与生活方式 ----
    ws2 = wb["家庭成员与生活方式"]
    key_map_2 = {
        "常住人口": "residents", "成员构成": "composition", "儿童年龄": "children_ages",
        "老人同住": "elderly", "宠物": "pets", "在家办公": "work_from_home",
        "待客频率": "entertain_frequency", "做饭频率": "cooking_frequency",
        "用餐习惯": "dining_habit", "动线": "movement_preference",
        "收纳强度": "storage_need", "爱好": "hobbies",
    }
    family = FamilyInfo()
    for r in range(6, 18):
        label = ws2.cell(row=r, column=2).value
        value = ws2.cell(row=r, column=3).value
        if not label:
            continue
        val = value if value and not str(value).startswith("（") else ""
        for cn, en in key_map_2.items():
            if cn in str(label):
                setattr(family, en, val)
                break
        else:
            setattr(family, str(label), val)
    conditions.family = family

    # ---- Sheet 3: 风格偏好 ----
    ws3 = wb["风格偏好"]
    key_map_3 = {
        "整体风格": "primary_style", "色调倾向": "color_tone",
        "设计关键词": "keywords", "地面材质": "floor_material",
        "墙面材质": "wall_material", "天花造型": "ceiling_type",
    }
    style = StylePreference()
    for r in range(6, 14):
        label = ws3.cell(row=r, column=2).value
        value = ws3.cell(row=r, column=3).value
        if not label:
            continue
        val = value if value and not str(value).startswith("（") else ""
        for cn, en in key_map_3.items():
            if cn in str(label):
                setattr(style, en, val)
                break
        else:
            setattr(style, str(label), val)
    conditions.style = style

    # ---- Sheet 4: 各空间需求 ----
    if "各空间需求" in [ws.title for ws in wb.worksheets]:
        ws4 = wb["各空间需求"]
        for r in range(6, ws4.max_row + 1):
            room_name = ws4.cell(row=r, column=2).value
            req = ws4.cell(row=r, column=3).value
            if room_name and req:
                rname = str(room_name)
                found = False
                for rm in conditions.rooms:
                    if rm["name"] == rname:
                        rm["requirements"][rname] = str(req)
                        found = True
                        break
                if not found:
                    conditions.rooms.append({"name": rname, "requirements": {rname: str(req)}})

    # ---- Sheet 5: 预算 ----
    if "预算" in [ws.title for ws in wb.worksheets]:
        ws5 = wb["预算"]
        budget = BudgetInfo()
        key_map_5 = {"总预算": "total_budget", "硬装": "hard_decoration",
                     "软装": "soft_decoration", "电器": "appliances", "备注": "notes"}
        for r in range(6, 12):
            label = ws5.cell(row=r, column=2).value
            value = ws5.cell(row=r, column=3).value
            if not label:
                continue
            for cn, en in key_map_5.items():
                if cn in str(label):
                    try:
                        setattr(budget, en, float(str(value).replace("万", "").replace(",", "")))
                    except (ValueError, TypeError):
                        setattr(budget, en, str(value) if value else "")
                    break
        conditions.budget = budget

    # ---- Sheet 6: 特殊需求 ----
    if "特殊需求" in [ws.title for ws in wb.worksheets]:
        ws6 = wb["特殊需求"]
        for r in range(6, ws6.max_row + 1):
            label = ws6.cell(row=r, column=2).value
            value = ws6.cell(row=r, column=3).value
            if label and value:
                conditions.special_requirements[str(label)] = str(value)

    conditions.source_note = f"Parsed from: {filepath}"
    return conditions


def to_dict(conditions: DesignConditions) -> dict:
    """将 DesignConditions 转为可序列化的字典"""
    return {
        "project": conditions.project.__dict__,
        "family": conditions.family.__dict__,
        "style": conditions.style.__dict__,
        "rooms": conditions.rooms,
        "budget": conditions.budget.__dict__,
        "special_requirements": conditions.special_requirements,
        "source_note": conditions.source_note,
    }


def save_conditions(conditions: DesignConditions, output_path: str):
    """保存设计条件为 JSON 文件"""
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(to_dict(conditions), f, ensure_ascii=False, indent=2)
    return output_path
