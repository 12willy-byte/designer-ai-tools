"""
模板生成器 - 生成设计需求问卷 Excel 模板
"""
import os
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def generate_template(output_path: str) -> str:
    wb = Workbook()

    header_font = Font(name="微软雅黑", bold=True, size=12, color="FFFFFF")
    bold_font = Font(name="微软雅黑", bold=True, size=10)
    normal_font = Font(name="微软雅黑", size=10)
    light_font = Font(name="微软雅黑", size=9, color="888888")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="B0B0B0"), right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"), bottom=Side(style="thin", color="B0B0B0"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = center_align; cell.border = thin_border

    def _set_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _add_title(ws, title, ncols=1):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        cell = ws.cell(row=1, column=1, value=title)
        cell.font = Font(name="微软雅黑", bold=True, size=14, color="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def _add_dv(ws, options, col):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.showDropDown = False; dv.prompt = "请选择"
        ws.add_data_validation(dv)
        return dv

    # ===== Sheet 1: 项目基本信息 =====
    ws1 = wb.active; ws1.title = "项目基本信息"
    ws1.sheet_properties.tabColor = "1F4E79"
    _set_widths(ws1, [5, 18, 30])
    _add_title(ws1, "室内设计需求调研表", 3)
    ws1.merge_cells("A3:C3")
    ws1.cell(row=3, column=1, value="填写说明：请在对应单元格中填写，单选题请选择选项").font = light_font
    row = 5
    for col, h in enumerate(["序号", "项目", "填写内容"], 1):
        ws1.cell(row=row, column=col, value=h)
    _style_header(ws1, row, 3)

    items1 = [
        ("1", "项目名称", ""), ("2", "项目地址", ""),
        ("3", "房屋类型", ["平层", "跃层", "别墅", "Loft", "工装"]),
        ("4", "建筑面积（m2）", ""), ("5", "原始户型图路径", ""),
        ("6", "设计类型", ["全案设计", "软装设计", "局部改造", "纯施工图"]),
        ("7", "期望完成时间", ""), ("8", "设计团队/设计师", ""),
    ]
    for i, (seq, label, val) in enumerate(items1):
        r = 6 + i
        ws1.cell(row=r, column=1, value=seq).font = normal_font
        ws1.cell(row=r, column=1).alignment = center_align
        ws1.cell(row=r, column=2, value=label).font = bold_font
        ws1.cell(row=r, column=2).alignment = left_align
        cell = ws1.cell(row=r, column=3)
        if isinstance(val, list):
            _add_dv(ws1, val, 3).add(cell)
        else:
            cell.value = val
        cell.font = normal_font; cell.alignment = left_align
        for c in range(1, 4):
            ws1.cell(row=r, column=c).border = thin_border

    # ===== Sheet 2: 家庭成员与生活方式 =====
    ws2 = wb.create_sheet("家庭成员与生活方式")
    ws2.sheet_properties.tabColor = "2E75B6"
    _set_widths(ws2, [5, 22, 30])
    _add_title(ws2, "家庭成员与生活方式", 3)
    row = 5
    for col, h in enumerate(["序号", "项目", "填写内容"], 1):
        ws2.cell(row=row, column=col, value=h)
    _style_header(ws2, row, 3)

    items2 = [
        ("1", "常住人口", ["1人", "2人", "3人", "4人", "5人及以上"]),
        ("2", "成员构成", ""), ("3", "儿童年龄", ""),
        ("4", "是否与老人同住", ["是", "否"]),
        ("5", "宠物", ["无", "猫", "狗", "其他"]),
        ("6", "是否在家办公", ["是，需要独立书房", "偶尔，餐桌/客厅即可", "否"]),
        ("7", "待客频率", ["很少", "偶尔", "经常", "频繁"]),
        ("8", "做饭频率", ["很少", "偶尔", "每天", "热爱烹饪"]),
        ("9", "用餐习惯", ["餐桌", "吧台", "岛台", "灵活多样"]),
        ("10", "动线偏好", ""), ("11", "收纳强度需求", ["一般", "较强", "极强"]),
        ("12", "兴趣爱好", ""),
    ]
    for i, (seq, label, val) in enumerate(items2):
        r = 6 + i
        ws2.cell(row=r, column=1, value=seq).font = normal_font
        ws2.cell(row=r, column=1).alignment = center_align
        ws2.cell(row=r, column=2, value=label).font = bold_font
        ws2.cell(row=r, column=2).alignment = left_align
        cell = ws2.cell(row=r, column=3)
        if isinstance(val, list):
            _add_dv(ws2, val, 3).add(cell)
        else:
            cell.value = val
        cell.font = normal_font; cell.alignment = left_align
        for c in range(1, 4):
            ws2.cell(row=r, column=c).border = thin_border

    # ===== Sheet 3: 风格偏好 =====
    ws3 = wb.create_sheet("风格偏好")
    ws3.sheet_properties.tabColor = "548235"
    _set_widths(ws3, [5, 22, 30])
    _add_title(ws3, "风格偏好", 3)
    row = 5
    for col, h in enumerate(["序号", "项目", "填写内容"], 1):
        ws3.cell(row=row, column=col, value=h)
    _style_header(ws3, row, 3)

    items3 = [
        ("1", "整体风格偏好", ["现代简约", "新中式", "北欧", "日式", "轻法式", "工业风", "美式", "混搭", "未确定"]),
        ("2", "色调倾向", ["暖色调", "冷色调", "中性色", "未确定"]),
        ("3", "设计关键词（逗号分隔）", ""),
        ("4", "地面材质偏好", ["木地板", "瓷砖", "大理石", "微水泥", "未确定"]),
        ("5", "墙面材质偏好", ["乳胶漆", "墙布", "木饰面", "微水泥", "未确定"]),
        ("6", "天花造型偏好", ["平顶", "边吊", "悬浮吊顶", "未确定"]),
        ("7", "参考图片/链接", ""),
    ]
    for i, (seq, label, val) in enumerate(items3):
        r = 6 + i
        ws3.cell(row=r, column=1, value=seq).font = normal_font
        ws3.cell(row=r, column=1).alignment = center_align
        ws3.cell(row=r, column=2, value=label).font = bold_font
        ws3.cell(row=r, column=2).alignment = left_align
        cell = ws3.cell(row=r, column=3)
        if isinstance(val, list):
            _add_dv(ws3, val, 3).add(cell)
        else:
            cell.value = val
        cell.font = normal_font; cell.alignment = left_align
        for c in range(1, 4):
            ws3.cell(row=r, column=c).border = thin_border

    # ===== Sheet 4: 各空间需求 =====
    ws4 = wb.create_sheet("各空间需求")
    ws4.sheet_properties.tabColor = "BF8F00"
    _set_widths(ws4, [5, 18, 40])
    _add_title(ws4, "各空间设计需求", 3)
    ws4.merge_cells("A3:C3")
    ws4.cell(row=3, column=1, value="请列出每个空间的具体需求（如：主卧需要衣帽间、阳台需要茶室等）").font = light_font
    row = 5
    for col, h in enumerate(["序号", "空间名称", "需求描述"], 1):
        ws4.cell(row=row, column=col, value=h)
    _style_header(ws4, row, 3)
    rooms_sample = [("1", "客厅"), ("2", "餐厅"), ("3", "主卧"), ("4", "次卧/儿童房"),
                    ("5", "书房"), ("6", "厨房"), ("7", "卫生间"), ("8", "阳台"), ("9", "玄关")]
    for i, (seq, name) in enumerate(rooms_sample):
        r = 6 + i
        ws4.cell(row=r, column=1, value=seq).font = normal_font
        ws4.cell(row=r, column=1).alignment = center_align
        ws4.cell(row=r, column=2, value=name).font = bold_font
        ws4.cell(row=r, column=3).font = normal_font
        for c in range(1, 4):
            ws4.cell(row=r, column=c).border = thin_border

    # ===== Sheet 5: 预算 =====
    ws5 = wb.create_sheet("预算")
    ws5.sheet_properties.tabColor = "C00000"
    _set_widths(ws5, [5, 22, 30])
    _add_title(ws5, "预算信息", 3)
    row = 5
    for col, h in enumerate(["序号", "项目", "填写内容"], 1):
        ws5.cell(row=row, column=col, value=h)
    _style_header(ws5, row, 3)
    items5 = [("1", "总预算（万元）"), ("2", "硬装预算（万元）"),
              ("3", "软装预算（万元）"), ("4", "电器预算（万元）"), ("5", "备注")]
    for i, (seq, label) in enumerate(items5):
        r = 6 + i
        ws5.cell(row=r, column=1, value=seq).font = normal_font
        ws5.cell(row=r, column=1).alignment = center_align
        ws5.cell(row=r, column=2, value=label).font = bold_font
        ws5.cell(row=r, column=2).alignment = left_align
        ws5.cell(row=r, column=3).font = normal_font
        for c in range(1, 4):
            ws5.cell(row=r, column=c).border = thin_border

    # ===== Sheet 6: 特殊需求 =====
    ws6 = wb.create_sheet("特殊需求")
    ws6.sheet_properties.tabColor = "7030A0"
    _set_widths(ws6, [5, 22, 40])
    _add_title(ws6, "特殊需求与其他", 3)
    row = 5
    for col, h in enumerate(["序号", "项目", "填写内容"], 1):
        ws6.cell(row=row, column=col, value=h)
    _style_header(ws6, row, 3)
    items6 = [
        ("1", "环保等级要求", ["国标E1", "欧标E0", "ENF", "无特别要求"]),
        ("2", "智能家居需求", ["全屋智能", "基础智能（灯控+窗帘）", "不需要"]),
        ("3", "是否需要全屋定制", ["是", "否", "部分空间"]),
        ("4", "地暖/新风/中央空调需求", ""), ("5", "其他特殊需求", ""),
    ]
    for i, (seq, label, val) in enumerate(items6):
        r = 6 + i
        ws6.cell(row=r, column=1, value=seq).font = normal_font
        ws6.cell(row=r, column=1).alignment = center_align
        ws6.cell(row=r, column=2, value=label).font = bold_font
        ws6.cell(row=r, column=2).alignment = left_align
        cell = ws6.cell(row=r, column=3)
        if isinstance(val, list):
            _add_dv(ws6, val, 3).add(cell)
        else:
            cell.value = val
        cell.font = normal_font; cell.alignment = left_align
        for c in range(1, 4):
            ws6.cell(row=r, column=c).border = thin_border

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    return output_path
