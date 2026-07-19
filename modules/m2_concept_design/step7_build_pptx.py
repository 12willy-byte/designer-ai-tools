"""Step 7: 概念方案 PPT 打包器
把前面所有产出合成一份可演示的 PPT
"""
import json, os, sys, math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
from core.design_schema import normalize_conditions


def build_pptx(conditions_json_path, output_path):
    with open(conditions_json_path, "r", encoding="utf-8") as f:
        conditions = normalize_conditions(json.load(f))

    base_dir = os.path.dirname(os.path.abspath(conditions_json_path))
    out_dir = os.path.join(base_dir, "concept_output")
    os.makedirs(out_dir, exist_ok=True)

    project = conditions.get("project", {})
    style = conditions.get("style", {})
    space = conditions.get("space_data", {})
    brief_path = os.path.join(out_dir, "设计定位.txt")

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    W = prs.slide_width
    H = prs.slide_height

    COLORS = {
        "dark": RGBColor(0x1A, 0x1A, 0x2E),
        "accent": RGBColor(0xC4, 0xA4, 0x84),
        "gray": RGBColor(0x66, 0x66, 0x66),
        "light": RGBColor(0xF5, 0xF3, 0xF0),
        "white": RGBColor(0xFF, 0xFF, 0xFF),
    }

    def add_bg(slide, color=COLORS["light"]):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_textbox(slide, left, top, width, height, text, font_size=18,
                    color=COLORS["dark"], bold=False, align=PP_ALIGN.LEFT):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.alignment = align
        return txBox

    def add_img(slide, img_path, left, top, width, height=None):
        if not os.path.exists(img_path):
            return None
        if height is None:
            height = width * 0.75
        return slide.shapes.add_picture(img_path, left, top, width, height)

    def add_rect(slide, left, top, width, height, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return shape

    # ── Slide 1: 封面 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_bg(slide, COLORS["dark"])

    # 装饰条
    add_rect(slide, 0, Inches(3.2), W, Pt(4), COLORS["accent"])

    project_name = project.get("name", "设计项目")
    add_textbox(slide, Inches(1), Inches(1.5), Inches(10), Inches(1.5),
               project_name, font_size=48, color=COLORS["white"], bold=True)

    style_name = style.get("primary_style", "现代简约")
    add_textbox(slide, Inches(1), Inches(3.6), Inches(10), Inches(0.8),
               "%s  ·  概念方案" % style_name, font_size=28, color=COLORS["accent"])

    area_str = ""
    building_area = project.get("area_m2") or project.get("area")
    if building_area:
        area_str += "建筑面积 %s m²" % building_area
    elif space.get("total_area_m2"):
        area_str += "房间加总约 %.0f m²（不含公摊/墙体）" % space["total_area_m2"]
    if space.get("total_rooms"):
        area_str += "  ·  %d 个房间" % space["total_rooms"]
    add_textbox(slide, Inches(1), Inches(4.6), Inches(10), Inches(0.6),
               area_str, font_size=18, color=RGBColor(0x88, 0x88, 0x88))

    add_textbox(slide, Inches(1), Inches(6.5), Inches(10), Inches(0.5),
               "AI 辅助概念设计方案", font_size=14, color=RGBColor(0x66, 0x66, 0x66))

    # ── Slide 2: 设计定位 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)

    add_rect(slide, 0, 0, Pt(8), H, COLORS["accent"])
    add_textbox(slide, Inches(0.6), Inches(0.5), Inches(1), Inches(1),
               "01", font_size=48, color=COLORS["accent"], bold=True)
    add_textbox(slide, Inches(0.6), Inches(1.3), Inches(2), Inches(0.5),
               "设计定位", font_size=24, color=COLORS["dark"], bold=True)

    brief_text = ""
    if os.path.exists(brief_path):
        brief_text = open(brief_path, "r", encoding="utf-8").read()
    if not brief_text:
        brief_text = "以中性色为基调的现代简约住宅，通过木饰面与微水泥的材质碰撞，营造温润有序的居住空间。"

    add_textbox(slide, Inches(2.5), Inches(2), Inches(9), Inches(4),
               brief_text, font_size=22, color=COLORS["gray"])

    # 关键词标签
    keywords = style.get("keywords", "现代,简约,自然")
    kws = [k.strip() for k in keywords.replace("、", " ").split() if k.strip()]
    if kws:
        x_pos = Inches(2.5)
        for kw in kws[:5]:
            add_rect(slide, x_pos, Inches(5.5), Inches(1.5), Inches(0.45), COLORS["accent"])
            add_textbox(slide, x_pos, Inches(5.5), Inches(1.5), Inches(0.45),
                       kw, font_size=16, color=COLORS["white"], align=PP_ALIGN.CENTER)
            x_pos += Inches(1.7)

    # ── Slide 3: 空间概况 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)

    add_textbox(slide, Inches(0.8), Inches(0.4), Inches(2), Inches(0.6),
               "02  空间概况", font_size=28, color=COLORS["dark"], bold=True)

    if space.get("rooms"):
        y = Inches(1.5)
        for rm in space["rooms"]:
            # 房间名
            add_rect(slide, Inches(0.8), y, Inches(2.5), Inches(0.6), COLORS["accent"])
            add_textbox(slide, Inches(0.8), y, Inches(2.5), Inches(0.6),
                       rm["name"], font_size=18, color=COLORS["white"], align=PP_ALIGN.CENTER)
            # 尺寸
            add_textbox(slide, Inches(3.6), y, Inches(3), Inches(0.6),
                       "%d×%d mm" % (rm["width_mm"], rm["height_mm"]),
                       font_size=16, color=COLORS["dark"])
            # 面积
            add_textbox(slide, Inches(6.5), y, Inches(2), Inches(0.6),
                       "%.1f m²" % rm["area_m2"], font_size=16, color=COLORS["gray"])
            y += Inches(0.8)

    # ── Slide 4: 色彩方案 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)

    add_textbox(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.6),
               "03  色彩方案", font_size=28, color=COLORS["dark"], bold=True)

    palette_img = os.path.join(out_dir, "色彩方案.png")
    if os.path.exists(palette_img):
        add_img(slide, palette_img, Inches(0.8), Inches(1.3), Inches(11))

    # ── Slide 5: 材质方案 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)

    add_textbox(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.6),
               "04  材质方案", font_size=28, color=COLORS["dark"], bold=True)

    mat_img = os.path.join(out_dir, "材质方案.png")
    if os.path.exists(mat_img):
        add_img(slide, mat_img, Inches(0.8), Inches(1.3), Inches(11))

    # ── Slide 6: 风格意向板 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)

    add_textbox(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.6),
               "05  风格意向板", font_size=28, color=COLORS["dark"], bold=True)

    mood_img = os.path.join(out_dir, "风格意向板.png")
    if os.path.exists(mood_img):
        add_img(slide, mood_img, Inches(0.8), Inches(1.2), Inches(11.5))

    # ── Slide 7: 布局方案 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)

    add_textbox(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.6),
               "06  布局方案", font_size=28, color=COLORS["dark"], bold=True)

    # 布局方案的文字摘要（优先读 M4 布局草案 JSON；无 JSON 时回退旧的 DXF 文字）
    layout_json_path = os.path.join(out_dir, "布局方案.json")
    layout_data = None
    if os.path.exists(layout_json_path):
        try:
            layout_data = json.load(open(layout_json_path, "r", encoding="utf-8"))
        except Exception:
            layout_data = None

    if layout_data and layout_data.get("status") == "draft" and layout_data.get("rooms"):
        rooms = layout_data["rooms"]
        add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11.5), Inches(0.5),
                   "布局草案（M4）覆盖 %d 个空间，供讨论复核，非施工依据" % len(rooms),
                   font_size=15, color=COLORS["gray"])
        y = Inches(1.7)
        for rm in rooms:
            zones = "、".join(z.get("name", "") for z in rm.get("zones") or [])
            add_textbox(slide, Inches(0.9), y, Inches(11.5), Inches(0.4),
                       "%s：%s" % (rm.get("name", ""), zones or "功能分区待确认"),
                       font_size=15, color=COLORS["dark"], bold=True)
            y += Inches(0.42)
            furniture = rm.get("furniture") or []
            if furniture:
                items = "；".join(
                    "%s %dx%dmm" % (f.get("item", ""), f.get("width_mm", 0), f.get("depth_mm", 0))
                    for f in furniture[:4])
                add_textbox(slide, Inches(1.2), y, Inches(11), Inches(0.4),
                           items, font_size=13, color=COLORS["gray"])
                y += Inches(0.4)
            confirms = rm.get("confirm_points") or []
            if confirms:
                add_textbox(slide, Inches(1.2), y, Inches(11), Inches(0.4),
                           "现场确认：%s" % confirms[0], font_size=12, color=COLORS["accent"])
                y += Inches(0.38)
            if y > Inches(6.6):
                break
    elif layout_data and layout_data.get("status", "").startswith("blocked"):
        reasons = layout_data.get("reasons") or []
        add_textbox(slide, Inches(0.8), Inches(1.4), Inches(11.5), Inches(0.6),
                   "布局草案未生成：空间事实不足，自动化闸门已安全拦截（未伪造布局）。",
                   font_size=16, color=COLORS["dark"], bold=True)
        y = Inches(2.2)
        for reason in reasons[:6]:
            add_textbox(slide, Inches(1), y, Inches(11), Inches(0.45),
                       "· " + reason, font_size=14, color=COLORS["gray"])
            y += Inches(0.45)
    else:
        layout_path = os.path.join(out_dir, "布局方案.dxf")
        if os.path.exists(layout_path):
            # 读取 DXf 中的文字建议作为布局说明
            import ezdxf
            doc = ezdxf.readfile(layout_path)
            texts = []
            for e in doc.modelspace():
                if e.dxftype() == "TEXT" and e.dxf.layer.startswith("AI方案"):
                    texts.append(e.dxf.text)

            if texts:
                y = Inches(1.5)
                for t in texts[:15]:
                    add_textbox(slide, Inches(1), y, Inches(11), Inches(0.4),
                               t, font_size=14, color=COLORS["gray"])
                    y += Inches(0.38)
                    if y > Inches(6.5):
                        break

    # ── Slide: 预算总览（M5，闸门放行且已生成估算时出现）──
    budget_json_path = os.path.join(out_dir, "budget_estimate.json")
    budget_data = None
    if os.path.exists(budget_json_path):
        try:
            budget_data = json.load(open(budget_json_path, "r", encoding="utf-8"))
        except Exception:
            budget_data = None

    atmos_start = 7
    if budget_data and budget_data.get("status") == "estimate" and budget_data.get("rooms"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide)

        add_textbox(slide, Inches(0.8), Inches(0.4), Inches(6), Inches(0.6),
                   "07  预算总览", font_size=28, color=COLORS["dark"], bold=True)
        atmos_start = 8

        totals = budget_data.get("totals") or {}
        comparison = budget_data.get("budget_comparison") or {}
        tier = budget_data.get("tier") or {}
        add_textbox(slide, Inches(0.8), Inches(1.05), Inches(11.5), Inches(0.5),
                   "总价区间：%.0f ~ %.0f 元（约 %.1f ~ %.1f 万元） · 价格档次：%s（参考价位档，需按当地市场校准）" % (
                       totals.get("low", 0), totals.get("high", 0),
                       totals.get("low", 0) / 10000, totals.get("high", 0) / 10000,
                       tier.get("level", "")),
                   font_size=17, color=COLORS["dark"], bold=True)
        if comparison.get("conclusion"):
            add_textbox(slide, Inches(0.8), Inches(1.55), Inches(11.5), Inches(0.5),
                       "预算比对：%s" % comparison["conclusion"],
                       font_size=14, color=COLORS["accent"], bold=True)

        # 分房间汇总表
        y = Inches(2.25)
        add_rect(slide, Inches(0.8), y, Inches(2.2), Inches(0.45), COLORS["accent"])
        add_textbox(slide, Inches(0.8), y, Inches(2.2), Inches(0.45),
                   "房间", font_size=14, color=COLORS["white"], align=PP_ALIGN.CENTER)
        add_textbox(slide, Inches(3.2), y, Inches(1.8), Inches(0.45),
                   "面积(㎡)", font_size=14, color=COLORS["dark"], align=PP_ALIGN.CENTER)
        add_textbox(slide, Inches(5.2), y, Inches(3.2), Inches(0.45),
                   "小计区间(元)", font_size=14, color=COLORS["dark"], align=PP_ALIGN.CENTER)
        add_textbox(slide, Inches(8.6), y, Inches(3.6), Inches(0.45),
                   "主要分项", font_size=14, color=COLORS["dark"])
        y += Inches(0.5)
        for rm in budget_data["rooms"]:
            add_textbox(slide, Inches(0.8), y, Inches(2.2), Inches(0.4),
                       rm.get("name", ""), font_size=14, color=COLORS["dark"], bold=True,
                       align=PP_ALIGN.CENTER)
            add_textbox(slide, Inches(3.2), y, Inches(1.8), Inches(0.4),
                       "%.2f" % rm.get("area_m2", 0), font_size=13, color=COLORS["gray"],
                       align=PP_ALIGN.CENTER)
            add_textbox(slide, Inches(5.2), y, Inches(3.2), Inches(0.4),
                       "%.0f ~ %.0f" % (rm.get("subtotal_low", 0), rm.get("subtotal_high", 0)),
                       font_size=13, color=COLORS["gray"], align=PP_ALIGN.CENTER)
            top_items = sorted(rm.get("items") or [], key=lambda i: -i.get("subtotal_high", 0))[:2]
            add_textbox(slide, Inches(8.6), y, Inches(3.8), Inches(0.4),
                       "、".join(i.get("item", "") for i in top_items),
                       font_size=12, color=COLORS["gray"])
            y += Inches(0.44)
            if y > Inches(6.1):
                break
        whole = budget_data.get("whole_house_items") or []
        if whole and y <= Inches(6.1):
            add_textbox(slide, Inches(0.8), y, Inches(2.2), Inches(0.4),
                       "全屋项目", font_size=14, color=COLORS["dark"], bold=True,
                       align=PP_ALIGN.CENTER)
            add_textbox(slide, Inches(3.2), y, Inches(1.8), Inches(0.4), "—",
                       font_size=13, color=COLORS["gray"], align=PP_ALIGN.CENTER)
            add_textbox(slide, Inches(5.2), y, Inches(3.2), Inches(0.4),
                       "%.0f ~ %.0f" % (sum(i.get("subtotal_low", 0) for i in whole),
                                        sum(i.get("subtotal_high", 0) for i in whole)),
                       font_size=13, color=COLORS["gray"], align=PP_ALIGN.CENTER)
            add_textbox(slide, Inches(8.6), y, Inches(3.8), Inches(0.4),
                       "、".join(i.get("item", "") for i in whole[:3]),
                       font_size=12, color=COLORS["gray"])
            y += Inches(0.44)
        add_textbox(slide, Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.4),
                   "工程量与单价依据详见 budget_estimate.json / material_list.json；区间为 AI 参考估算，非报价单",
                   font_size=12, color=COLORS["gray"])

    # ── Slide 8~: 氛围图 ──
    # 找到所有氛围图
    atmos_files = sorted([f for f in os.listdir(out_dir) if "氛围图" in f and f.endswith(".png")])

    for i, atmos_file in enumerate(atmos_files):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide, COLORS["dark"])

        room_name = atmos_file.replace("氛围图.png", "")
        add_textbox(slide, Inches(0.8), Inches(0.3), Inches(6), Inches(0.6),
                   "0%s  %s 氛围图" % (str(i + atmos_start), room_name),
                   font_size=24, color=COLORS["accent"], bold=True)

        atmos_path = os.path.join(out_dir, atmos_file)
        if os.path.exists(atmos_path):
            # 图片居中
            img_w = Inches(9.5)
            img_h = Inches(5.5)
            left = (W - img_w) // 2
            top = Inches(1.3)
            add_img(slide, atmos_path, left, top, img_w, img_h)

    # ── 倒数第二页: 假设与口径说明 ──
    # 汇总各步骤的推断标注与面积口径，避免假设与事实混排流入交付物
    assumption_items = []

    # 1) 设计定位正文中以「假设」开头的行
    if brief_text:
        for line in brief_text.splitlines():
            stripped = line.strip().lstrip("-•·*").strip()
            if stripped.startswith("假设") and stripped not in assumption_items:
                assumption_items.append(stripped)

    # 2) 材质方案 JSON 的 assumptions 数组
    mat_json_path = os.path.join(out_dir, "材质方案.json")
    if os.path.exists(mat_json_path):
        try:
            mat_data = json.load(open(mat_json_path, "r", encoding="utf-8"))
            for item in mat_data.get("assumptions") or []:
                text = str(item).strip()
                if text and text not in assumption_items:
                    assumption_items.append(text)
        except Exception:
            pass

    # 3) 布局草案 JSON 的 assumptions 数组（M4）
    if layout_data and layout_data.get("status") == "draft":
        for item in layout_data.get("assumptions") or []:
            text = str(item).strip()
            if text and text not in assumption_items:
                assumption_items.append(text)

    # 3.5) 预算估算 JSON 的 assumptions 数组（M5）
    if budget_data and budget_data.get("status") == "estimate":
        for item in budget_data.get("assumptions") or []:
            text = str(item).strip()
            if text and text not in assumption_items:
                assumption_items.append(text)

    # 4) 面积口径说明（建筑面积 vs 房间加总）
    caliber_notes = []
    building_area = project.get("area_m2") or project.get("area")
    room_total = space.get("total_area_m2")
    if building_area and room_total:
        caliber_notes.append(
            "面积口径：本方案总面积以建筑面积 %s㎡ 为准；各房间加总约 %.1f㎡（不含公摊/墙体），两者差异属正常口径差。"
            % (building_area, room_total))
    elif room_total and not building_area:
        caliber_notes.append(
            "面积口径：输入未提供建筑面积，本方案面积均为房间加总约 %.1f㎡（不含公摊/墙体）。"
            % room_total)

    if assumption_items or caliber_notes:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide)

        add_rect(slide, 0, 0, Pt(8), H, COLORS["accent"])
        add_textbox(slide, Inches(0.6), Inches(0.5), Inches(6), Inches(0.7),
                   "假设与口径说明", font_size=28, color=COLORS["dark"], bold=True)
        add_textbox(slide, Inches(0.6), Inches(1.15), Inches(11), Inches(0.4),
                   "以下内容为 AI 推断或口径说明，非客户输入事实，请设计师复核后再向客户呈现",
                   font_size=13, color=COLORS["gray"])

        y = Inches(1.8)
        for note in caliber_notes:
            add_textbox(slide, Inches(0.8), y, Inches(11.5), Inches(0.5),
                       note, font_size=15, color=COLORS["dark"])
            y += Inches(0.55)
        for item in assumption_items[:16]:
            add_textbox(slide, Inches(0.8), y, Inches(11.5), Inches(0.45),
                       "· " + item, font_size=13, color=COLORS["gray"])
            y += Inches(0.42)
            if y > Inches(6.9):
                break

    # ── 最后一页: 结尾 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["dark"])

    add_rect(slide, Inches(4.5), Inches(3.5), Inches(4), Pt(3), COLORS["accent"])
    add_textbox(slide, Inches(1), Inches(2.5), Inches(11), Inches(1),
               "谢谢观看", font_size=44, color=COLORS["white"], bold=True,
               align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(4), Inches(11), Inches(0.6),
               "期待与您共同实现这个理想家", font_size=20, color=COLORS["accent"],
               align=PP_ALIGN.CENTER)

    # 保存
    prs.save(output_path)
    return output_path


if __name__ == "__main__":
    inp = sys.argv[1] if len(sys.argv) > 1 else "../templates/设计条件.json"
    out = sys.argv[2] if len(sys.argv) > 2 else "../templates/concept_output/概念方案.pptx"
    r = build_pptx(inp, out)
    sz = os.path.getsize(out)
    print("OK: %s (%d bytes, %d slides)" % (r, sz,
          len(Presentation(out).slides)))
