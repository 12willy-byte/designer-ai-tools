# ============================================================
# 模板生成器: 生成设计需求问卷 Excel 模板
# ============================================================
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def generate_template(output_path):
    """生成室内设计需求问卷 Excel 模板"""
    wb = Workbook()

    # ===== 样式 =====
    header_font = Font(name="微软雅黑", bold=True, size=12, color="FFFFFF")
    bold_font = Font(name="微软雅黑", bold=True, size=10)
    normal_font = Font(name="微软雅黑", size=10)
    light_font = Font(name="微软雅黑", size=9, color="888888")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

    def _set_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _add_title(ws, title, ncols=6):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        cell = ws.cell(row=1, column=1, value=title)
        cell.font = Font(name="微软雅黑", bold=True, size=14, color="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def _add_dv(ws, options):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.showDropDown = False
        dv.prompt = "请选择"
        ws.add_data_validation(dv)
        return dv

    # ---- Sheet 1: 项目基本信息 ----
    ws1 = wb.active
    ws1.title = "项目基本信息"
    ws1.sheet_properties.tabColor = "1F4E79"
    _set_widths(ws1, [5, 18, 30, 30, 5])
    _add_title(ws1, "室内设计需求调研表", 5)

    ws1.merge_cells("A3:E3")
    ws1.cell(row=3, column=1, value="填写说明：请在各对应单元格中填写内容，单选题请选择选项，多选题可填写多个（逗号分隔）").font = light_font

    row = 5
    for col, h in enumerate(["序号", "项目", "填写内容", "", ""], 1):
        ws1.cell(row=row, column=col, value=h)
    _style_header(ws1, row, 5)
    ws1.merge_cells("B5:C5")
    ws1.merge_cells("D5:E5")

    items = [
        ("1", "项目名称", ""),
        ("2", "项目地址", ""),
        ("3", "房屋类型", ["平层", "跃层", "别墅", "Loft", "工装"]),
        ("4", "建筑面积（m²）", ""),
        ("5", "原始户型图路径", ""),
        ("6", "设计类型", ["全案设计", "软装设计", "局部改造", "纯施工图"]),
        ("7", "期望完成时间", ""),
        ("8", "设计团队/设计师", ""),
    ]
    for i, (seq, label, value) in enumerate(items):
        r = 6 + i
        ws1.cell(row=r, column=1, value=seq).font = normal_font
        ws1.cell(row=r, column=1).alignment = center_align
        ws1.cell(row=r, column=2, value=label).font = bold_font
        ws1.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws1.cell(row=r, column=2).alignment = left_align
        ws1.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        cell = ws1.cell(row=r, column=4)
        if isinstance(value, list):
            cell.value = ""
            _add_dv(ws1, value).add(cell)
        else:
            cell.value = value
        cell.font = normal_font
        cell.alignment = left_align
        for c in range(1, 6):
            ws1.cell(row=r, column=c).border = thin_border

    # ---- Sheet 2: 家庭成员与生活方式 ----
    ws2 = wb.create_sheet("家庭成员与生活方式")
    ws2.sheet_properties.tabColor = "2E75B6"
    _set_widths(ws2, [5, 20, 35, 35, 5])
    _add_title(ws2, "家庭成员与生活方式", 5)

    ws2.merge_cells("A3:E3")
    ws2.cell(row=3, column=1, value="填写说明：单选题请从下拉菜单选择，多选可填多个（逗号分隔）").font = light_font

    row = 5
    for col, h in enumerate(["序号", "调查项", "填写内容/选项", "备注说明", ""], 1):
        ws2.cell(row=row, column=col, value=h)
    _style_header(ws2, row, 5)
    ws2.merge_cells("D5:E5")

    items = [
        ("1", "常住人口（人）", None, ""),
        ("2", "成员构成", ["独居", "夫妻", "夫妻+儿童", "三代同堂", "合租"], ""),
        ("3", "儿童年龄", None, "如有多个逗号分隔"),
        ("4", "有老人同住", ["否", "是（自理）", "是（需照料）"], ""),
        ("5", "宠物", ["无", "狗", "猫", "狗+猫", "其他"], ""),
        ("6", "在家办公需求", ["不需要", "需要独立书房", "需要灵活工位", "偶尔在家办公"], ""),
        ("7", "待客频率", ["极少", "偶尔", "经常聚会", "频繁社交"], ""),
        ("8", "做饭频率", ["很少做饭", "每日做饭", "喜欢烘焙", "中式爆炒为主", "轻食为主"], ""),
        ("9", "用餐习惯", ["在家用餐为主", "外卖为主", "周末才在家吃"], ""),
        ("10", "居家动线偏好/习惯", None, ""),
        ("11", "储物收纳强度", ["普通", "较多（有囤积习惯）", "极简（物品很少）"], ""),
        ("12", "爱好/收藏", None, ""),
    ]
    for i, (seq, label, options, note) in enumerate(items):
        r = 6 + i
        ws2.cell(row=r, column=1, value=seq).font = normal_font
        ws2.cell(row=r, column=1).alignment = center_align
        ws2.cell(row=r, column=2, value=label).font = bold_font
        ws2.cell(row=r, column=2).alignment = left_align
        ws2.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        nc = ws2.cell(row=r, column=4, value=note)
        nc.font = light_font
        cell = ws2.cell(row=r, column=3)
        cell.font = normal_font
        cell.alignment = left_align
        if options:
            _add_dv(ws2, options).add(cell)
        for c in range(1, 6):
            ws2.cell(row=r, column=c).border = thin_border

    # ---- Sheet 3: 风格偏好 ----
    ws3 = wb.create_sheet("风格偏好")
    ws3.sheet_properties.tabColor = "70AD47"
    _set_widths(ws3, [5, 18, 35, 35, 5])
    _add_title(ws3, "风格偏好", 5)

    row = 5
    for col, h in enumerate(["序号", "调查项", "填写内容/选项", "备注说明", ""], 1):
        ws3.cell(row=row, column=col, value=h)
    _style_header(ws3, row, 5)
    ws3.merge_cells("D5:E5")

    items = [
        ("1", "主设计风格", ["现代简约", "极简", "新中式", "日式/原木", "法式/轻法式", "美式", "工业风", "北欧", "复古/中古", "混搭", "奶油风", "侘寂风", "其他"]),
        ("2", "色调倾向", ["暖色调", "冷色调", "中性色/大地色", "黑白灰", "自然原木色", "大胆撞色", "奶油色系"]),
        ("3", "墙面材质偏好", ["乳胶漆", "墙纸/墙布", "艺术涂料", "木饰面", "石材/岩板", "微水泥", "护墙板"]),
        ("4", "地面材质偏好", ["实木地板", "复合地板", "瓷砖", "木纹砖", "微水泥", "石材", "地毯满铺"]),
        ("5", "天花板偏好", ["平顶（不吊顶）", "局部吊顶", "全吊顶无主灯", "石膏线装饰", "悬浮顶"]),
        ("6", "门的风格", ["平板门", "回字造型门", "玻璃门", "隐形门", "移门/谷仓门"]),
        ("7", "喜欢的参考图路径", None),
        ("8", "不喜欢的风格/元素", None),
        ("9", "设计关键词（3-5个）", None),
    ]
    for i, (seq, label, options) in enumerate(items):
        r = 6 + i
        ws3.cell(row=r, column=1, value=seq).font = normal_font
        ws3.cell(row=r, column=1).alignment = center_align
        ws3.cell(row=r, column=2, value=label).font = bold_font
        ws3.cell(row=r, column=2).alignment = left_align
        ws3.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        cell = ws3.cell(row=r, column=3)
        cell.font = normal_font
        cell.alignment = left_align
        if options:
            _add_dv(ws3, options).add(cell)
        for c in range(1, 6):
            ws3.cell(row=r, column=c).border = thin_border

    # ---- Sheet 4: 空间需求分解（每个房间细分） ----
    ws4 = wb.create_sheet("空间需求分解")
    ws4.sheet_properties.tabColor = "ED7D31"
    _set_widths(ws4, [4, 14, 28, 40, 5])
    _add_title(ws4, "空间需求分解 — 逐空间详细设计需求", 5)

    row = 5
    for col, h in enumerate(["序号", "空间", "设计项目", "需求选择/填写内容", ""], 1):
        ws4.cell(row=row, column=col, value=h)
    _style_header(ws4, row, 5)

    rooms = [
        ("玄关", [("玄关柜", ["到顶柜", "矮柜", "换鞋凳一体", "悬挑柜", "不需要"]),
                   ("全身镜", ["需要", "不需要"]),
                   ("感应灯", ["需要", "不需要"]),
                   ("挂衣区", ["开放挂衣", "柜内挂衣", "不需要"]),
                   ("地面材质", ["瓷砖", "地板", "石材", "拼花"]),
                   ("换鞋凳", ["固定式", "移动式", "不需要"])]),
        ("客厅", [("背景墙造型", ["简约乳胶漆", "石材/岩板", "木饰面", "收纳柜式", "投影幕布", "无造型"]),
                   ("电视尺寸预留（英寸）", None),
                   ("设备收纳", ["隐藏式机柜", "开放格", "不需要"]),
                   ("沙发类型", ["直排", "L型", "单人位+躺椅", "组合模块", "不需要"]),
                   ("茶几", ["需要", "不需要", "边几代替"]),
                   ("地毯", ["需要", "不需要"]),
                   ("窗帘形式", ["落地帘", "百叶", "纱帘+布帘", "梦幻帘", "罗马帘"]),
                   ("阳台门改造", ["保留", "打通", "折叠门", "推拉门"]),
                   ("整面收纳柜", ["需要", "不需要"]),
                   ("展示/陈设区", ["需要", "不需要"]),
                   ("主灯形式", ["吸顶灯", "吊灯", "无主灯（磁吸轨道）", "无主灯（筒射灯）"]),
                   ("特殊需求", None)]),
        ("餐厅", [("餐桌形式", ["圆桌", "长桌", "卡座+桌", "岛台+餐桌一体"]),
                   ("座位数", None),
                   ("餐边柜", ["到顶柜", "矮柜", "中空操作台", "玻璃柜门", "不需要"]),
                   ("内嵌电器", ["饮水机", "咖啡机", "管线机", "酒柜", "不需要"]),
                   ("主灯形式", ["吊线灯", "吸顶灯", "无主灯"]),
                   ("特殊需求", None)]),
        ("厨房", [("厨房布局", ["一字型", "L型", "U型", "双一字", "中岛", "G型"]),
                   ("烟机灶具", ["集成灶", "分体烟灶", "分体+消毒柜", "分体+蒸烤箱"]),
                   ("台面材质", ["石英石", "岩板", "不锈钢", "实木", "陶瓷"]),
                   ("挡水条", ["前挡水", "后挡水", "无挡水"]),
                   ("水槽形式", ["单槽", "双槽", "台下盆", "台上盆", "花岗岩水槽"]),
                   ("净水系统", ["不需要", "直饮", "全屋净水", "软水"]),
                   ("吊柜", ["满吊", "局部吊", "玻璃门吊柜", "不需要"]),
                   ("地柜抽屉", ["普通层板", "全拉抽屉", "碗碟拉篮", "调味拉篮"]),
                   ("高柜", ["需要（嵌入烤箱/微波炉）", "需要（嵌入式冰箱）", "不需要"]),
                   ("厨房电器清单", None),
                   ("特殊需求", None)]),
        ("主卧", [("功能定位", ["纯睡眠", "睡眠+梳妆", "睡眠+办公", "睡眠+休闲"]),
                   ("床尺寸", ["1.5m", "1.8m", "2.0m", "定制"]),
                   ("床头背景", ["乳胶漆", "硬包/软包", "木饰面", "无造型", "墙纸/墙布"]),
                   ("床头柜", ["两侧各一", "单侧", "床头柜+梳妆台一体", "不需要"]),
                   ("衣柜形式", ["开门柜", "移门柜", "步入式衣帽间", "L型衣帽间", "U型衣帽间"]),
                   ("柜门材质", ["平板门", "玻璃门", "烤漆", "肤感", "PET"]),
                   ("衣柜内部功能", ["挂衣为主", "叠放为主", "混合"]),
                   ("梳妆台", ["需要", "不需要"]),
                   ("书桌", ["需要", "不需要"]),
                   ("主灯形式", ["吸顶灯", "吊灯", "无主灯", "不需要"]),
                   ("床头阅读灯", ["壁灯", "台灯", "吊线灯", "不需要"]),
                   ("特殊需求", None)]),
        ("次卧", [("功能定位", ["儿童房", "老人房", "客房", "多功能房", "书房+客房"]),
                   ("床尺寸", ["1.2m", "1.5m", "1.8m", "上下铺", "子母床", "隐形床"]),
                   ("书桌", ["需要", "不需要"]),
                   ("衣柜", ["需要", "不需要"]),
                   ("特殊需求", None)]),
        ("书房/多功能房", [("书桌形式", ["一字长桌", "L型转角桌", "双人位", "电动升降桌"]),
                            ("电脑和设备", ["台式机", "笔记本", "双屏", "打印机", "不需要"]),
                            ("书架", ["开放书架", "封闭柜门", "玻璃柜门", "整墙书柜"]),
                            ("休息区", ["沙发床", "单人榻", "不需要"]),
                            ("展示区", ["手办展示", "书籍展示", "不需要"]),
                            ("特殊需求", None)]),
        ("主卫", [("台盆形式", ["台上盆", "台下盆", "一体盆", "双台盆", "岩板台面"]),
                   ("浴室柜", ["吊柜", "落地柜", "镜柜", "双镜柜"]),
                   ("马桶", ["普通马桶", "智能马桶", "壁挂马桶"]),
                   ("淋浴形式", ["一字型淋浴房", "钻石型淋浴房", "浴缸+淋浴", "纯浴缸", "纯淋浴"]),
                   ("顶喷花洒", ["需要", "不需要"]),
                   ("壁龛", ["需要", "不需要"]),
                   ("暖风机/浴霸", ["需要", "不需要"]),
                   ("特殊需求", None)]),
        ("公卫", [("功能形式", ["三分离", "干湿分离", "常规", "开放式洗手台"]),
                   ("马桶/蹲便", ["马桶", "蹲便", "马桶+小便器"]),
                   ("台盆形式", ["台上盆", "台下盆", "一体盆"]),
                   ("淋浴形式", ["一字型淋浴房", "钻石型淋浴房", "浴帘", "纯淋浴"]),
                   ("收纳", ["镜柜", "壁龛", "置物架", "不需要"]),
                   ("特殊需求", None)]),
        ("阳台", [("功能定位", ["纯生活阳台", "纯休闲阳台", "生活+休闲", "封进室内"]),
                   ("洗衣机/烘干机", ["并排放", "叠放", "单独洗衣机", "单独烘干机", "不需要"]),
                   ("家政柜", ["需要", "不需要"]),
                   ("洗手台", ["需要", "不需要"]),
                   ("晾晒方式", ["手动晾衣架", "电动晾衣架", "隐藏式晾衣架", "烘干机代替"]),
                   ("休闲功能", ["绿植区", "茶桌", "健身区", "书桌", "不需要"]),
                   ("封窗", ["需要封窗", "不封窗", "不确定"]),
                   ("特殊需求", None)]),
        ("衣帽间", [("形式", ["定制衣柜", "开放挂衣系统", "IKEA博阿克塞", "金属衣帽间"]),
                     ("长短挂衣分区", ["需要", "不需要"]),
                     ("饰品柜/首饰抽", ["需要", "不需要"]),
                     ("换衣镜", ["需要", "不需要"]),
                     ("换衣凳", ["需要", "不需要"]),
                     ("特殊需求", None)]),
        ("储物间/家政间", [("形式", ["开放货架", "定制柜", "不需要"]),
                            ("收纳内容", None),
                            ("是否需要吸尘器充电位", ["需要", "不需要"]),
                            ("是否需要扫地机器人位", ["需要", "不需要"]),
                            ("特殊需求", None)]),
    ]

    current_row = 6
    seq = 1
    room_fill = PatternFill(start_color="FFF2E6", end_color="FFF2E6", fill_type="solid")
    for room_name, items in rooms:
        # 房间标题行
        ws4.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        cell = ws4.cell(row=current_row, column=1, value=f"【{room_name}】")
        cell.font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
        cell.fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        for c in range(1, 6):
            ws4.cell(row=current_row, column=c).border = thin_border
            ws4.cell(row=current_row, column=c).fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
        current_row += 1

        for item_name, item_def in items:
            ws4.cell(row=current_row, column=1, value=seq).font = normal_font
            ws4.cell(row=current_row, column=1).alignment = center_align
            ws4.cell(row=current_row, column=2, value=room_name).font = bold_font
            ws4.cell(row=current_row, column=2).alignment = center_align
            ws4.cell(row=current_row, column=2).fill = room_fill
            ws4.cell(row=current_row, column=3, value=item_name).font = normal_font
            ws4.cell(row=current_row, column=3).alignment = left_align
            ws4.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=5)
            cell = ws4.cell(row=current_row, column=4)
            cell.font = normal_font
            cell.alignment = left_align
            if isinstance(item_def, list):
                _add_dv(ws4, item_def).add(cell)
            for c in range(1, 6):
                ws4.cell(row=current_row, column=c).border = thin_border
            seq += 1
            current_row += 1

    # ---- Sheet 5: 预算明细 ----
    ws5 = wb.create_sheet("预算明细")
    ws5.sheet_properties.tabColor = "FFC000"
    _set_widths(ws5, [5, 5, 22, 18, 30, 5])
    _add_title(ws5, "预算明细表", 6)

    row = 5
    for col, h in enumerate(["序号", "编号", "项目类别", "预估金额（元）", "备注/说明", ""], 1):
        ws5.cell(row=row, column=col, value=h)
    _style_header(ws5, row, 6)
    ws5.merge_cells("E5:F5")

    items = [
        ("A", "拆改工程", "墙体拆除、新建、垃圾清运等"),
        ("B", "水电工程", "水管、线管、强弱电改造"),
        ("C", "泥瓦工程", "防水、贴砖、找平"),
        ("D", "木工程/吊顶", "吊顶、背景墙基层"),
        ("E", "定制柜体", "橱柜、衣柜、玄关柜、餐边柜"),
        ("F", "油漆/涂料工程", "墙面、顶面乳胶漆、艺术涂料"),
        ("G", "主材-瓷砖/地板", "全屋地砖、木地板"),
        ("H", "主材-石材/岩板", "台面、背景墙等"),
        ("I", "主材-门窗", "房门、推拉门、封窗"),
        ("J", "卫浴/洁具", "马桶、花洒、台盆、浴室柜"),
        ("K", "灯具/开关插座", "全屋灯具、开关面板"),
        ("L", "家具", "沙发、床、餐桌椅、茶几等"),
        ("M", "窗帘/软饰", "窗帘、地毯、挂画、绿植"),
        ("N", "家电", "冰箱、洗衣机、烟灶、空调等"),
        ("O", "智能家居", "智能灯光、窗帘电机、安防等"),
        ("P", "其他费用", "设计费、管理费、杂费等"),
    ]
    for i, (code, name, note) in enumerate(items):
        r = 7 + i
        ws5.cell(row=r, column=1, value=i + 1).font = normal_font
        ws5.cell(row=r, column=1).alignment = center_align
        ws5.cell(row=r, column=2, value=code).font = bold_font
        ws5.cell(row=r, column=2).alignment = center_align
        ws5.cell(row=r, column=3, value=name).font = normal_font
        ws5.cell(row=r, column=3).alignment = left_align
        price_cell = ws5.cell(row=r, column=4)
        price_cell.font = normal_font
        price_cell.alignment = center_align
        price_cell.number_format = "#,##0"
        ws5.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
        ws5.cell(row=r, column=5, value=note).font = light_font
        for c in range(1, 7):
            ws5.cell(row=r, column=c).border = thin_border

    # 合计
    tr = 7 + len(items)
    ws5.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=3)
    ws5.cell(row=tr, column=1, value="合  计").font = Font(name="微软雅黑", bold=True, size=12, color="FFFFFF")
    for c in range(1, 7):
        ws5.cell(row=tr, column=c).fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        ws5.cell(row=tr, column=c).border = thin_border
        ws5.cell(row=tr, column=c).alignment = center_align
    ws5.cell(row=tr, column=4).number_format = "#,##0"

    # ---- Sheet 6: 特殊需求 ----
    ws6 = wb.create_sheet("特殊需求")
    ws6.sheet_properties.tabColor = "C00000"
    _set_widths(ws6, [5, 22, 35, 35, 5])
    _add_title(ws6, "特殊需求", 5)

    row = 5
    for col, h in enumerate(["序号", "需求类别", "选择/填写内容", "补充说明", ""], 1):
        ws6.cell(row=row, column=col, value=h)
    _style_header(ws6, row, 5)
    ws6.merge_cells("D5:E5")

    items = [
        ("1", "智能家居需求", ["不需要", "全屋智能", "部分智能（灯光）", "部分智能（窗帘）", "部分智能（安防）"]),
        ("2", "无障碍/适老化", ["无", "老人扶手", "轮椅通道", "无高差地面", "卫生间防滑"]),
        ("3", "儿童安全设计", ["无", "圆角处理", "防夹手", "窗户防护", "防触电插座", "低矮开关"]),
        ("4", "环保等级要求", ["普通", "高环保（儿童房/孕妇）", "全房环保标准"]),
        ("5", "风水要求", None),
        ("6", "室内绿植/花艺", ["不需要", "少量点缀", "喜欢室内花园", "需要水景"]),
        ("7", "特殊材质偏好", None),
        ("8", "施工时间要求", None),
        ("9", "其他补充需求", None),
    ]
    for i, (seq, label, options) in enumerate(items):
        r = 6 + i
        ws6.cell(row=r, column=1, value=seq).font = normal_font
        ws6.cell(row=r, column=1).alignment = center_align
        ws6.cell(row=r, column=2, value=label).font = bold_font
        ws6.cell(row=r, column=2).alignment = left_align
        ws6.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        cell = ws6.cell(row=r, column=3)
        cell.font = normal_font
        cell.alignment = left_align
        if options:
            _add_dv(ws6, options).add(cell)
        for c in range(1, 6):
            ws6.cell(row=r, column=c).border = thin_border

    # ---- Sheet 7: 客户签字确认 ----
    ws7 = wb.create_sheet("客户签字确认")
    ws7.sheet_properties.tabColor = "7030A0"
    _set_widths(ws7, [5, 20, 30, 30, 5])
    _add_title(ws7, "客户签字确认", 5)

    ws7.merge_cells("A5:E5")
    ws7.cell(row=5, column=1, value="本人确认上述需求信息真实、准确，同意设计师据此进行方案设计。").font = Font(name="微软雅黑", size=11)

    row = 7
    for col, h in enumerate(["序号", "项目", "填写内容", "", ""], 1):
        ws7.cell(row=row, column=col, value=h)
    _style_header(ws7, row, 5)

    for i, label in enumerate(["客户姓名", "联系电话", "签字日期", "备注"]):
        r = 8 + i
        ws7.cell(row=r, column=1, value=i + 1).font = normal_font
        ws7.cell(row=r, column=1).alignment = center_align
        ws7.cell(row=r, column=2, value=label).font = bold_font
        ws7.cell(row=r, column=2).alignment = left_align
        ws7.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        for c in range(1, 6):
            ws7.cell(row=r, column=c).border = thin_border

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else "设计需求问卷模板.xlsx"
    result = generate_template(path)
    print(f"OK: {result}")
