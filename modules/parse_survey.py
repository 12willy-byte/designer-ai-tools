# ============================================================
# 问卷解析器: Excel → JSON（AI 输入条件）
# ============================================================
import json
import openpyxl
from collections import OrderedDict


def parse_survey(filepath):
    """解析填好的设计需求问卷.xlsx，返回结构化字典"""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    result = OrderedDict()

    # ---- Sheet 1 ----
    ws1 = wb["项目基本信息"]
    project = OrderedDict()
    key_map_1 = {
        "项目名称": "name", "项目地址": "address", "房屋类型": "house_type",
        "建筑面积": "area", "户型图": "floor_plan_path", "设计类型": "design_type",
        "完成时间": "expected_completion", "设计师": "designer",
    }
    for r in range(6, 14):
        label = ws1.cell(row=r, column=2).value
        value = ws1.cell(row=r, column=4).value
        if not label:
            continue
        val = value if value and not str(value).startswith("（") else ""
        matched = False
        for cn, en in key_map_1.items():
            if cn in str(label):
                project[en] = val
                matched = True
                break
        if not matched:
            project[str(label)] = val
    result["project"] = project

    # ---- Sheet 2 ----
    ws2 = wb["家庭成员与生活方式"]
    family = OrderedDict()
    key_map_2 = {
        "常住人口": "residents", "成员构成": "composition", "儿童年龄": "children_ages",
        "老人同住": "elderly", "宠物": "pets", "在家办公": "work_from_home",
        "待客频率": "entertain_frequency", "做饭频率": "cooking_frequency",
        "用餐习惯": "dining_habit", "动线": "movement_preference",
        "收纳强度": "storage_need", "爱好": "hobbies",
    }
    for r in range(6, 18):
        label = ws2.cell(row=r, column=2).value
        value = ws2.cell(row=r, column=3).value
        if not label:
            continue
        val = value if value and not str(value).startswith("（") else ""
        for cn, en in key_map_2.items():
            if cn in str(label):
                family[en] = val
                break
        else:
            family[str(label)] = val
    result["family"] = family

    # ---- Sheet 3 ----
    ws3 = wb["风格偏好"]
    style = OrderedDict()
    key_map_3 = {
        "主设计风格": "primary_style", "色调": "color_tone", "墙面材质": "wall_material",
        "地面材质": "floor_material", "天花板": "ceiling_preference",
        "门的风格": "door_style", "参考图": "reference_images",
        "不喜欢": "dislikes", "关键词": "keywords",
    }
    for r in range(6, 15):
        label = ws3.cell(row=r, column=2).value
        value = ws3.cell(row=r, column=3).value
        if not label:
            continue
        val = value if value and not str(value).startswith("（") else ""
        for cn, en in key_map_3.items():
            if cn in str(label):
                style[en] = val
                break
        else:
            style[str(label)] = val
    result["style"] = style

    # ---- Sheet 4 ----
    ws4 = wb["空间需求分解"]
    rooms = OrderedDict()
    current_room = None
    for r in range(6, ws4.max_row + 1):
        v1 = ws4.cell(row=r, column=1).value
        v2 = ws4.cell(row=r, column=2).value
        v3 = ws4.cell(row=r, column=3).value
        v4 = ws4.cell(row=r, column=4).value

        if v1 and isinstance(v1, str) and v1.startswith("【"):
            current_room = v1.replace("【", "").replace("】", "").strip()
            rooms[current_room] = OrderedDict()
            continue

        if current_room and v3:
            val = v4 if v4 and not str(v4).startswith("（") else ""
            if val and "," in str(val):
                rooms[current_room][v3] = [x.strip() for x in str(val).split(",")]
            else:
                rooms[current_room][v3] = val
    result["rooms"] = rooms

    # ---- Sheet 5 ----
    ws5 = wb["预算明细"]
    budget = OrderedDict()
    budget_map = {
        "拆改": "demolition", "水电": "plumbing_electrical", "泥瓦": "tiling",
        "木工": "carpentry", "定制柜": "custom_cabinets", "油漆": "painting",
        "瓷砖": "flooring_tiles", "石材": "stone", "门窗": "doors_windows",
        "卫浴": "bathroom", "灯具": "lighting", "家具": "furniture",
        "窗帘": "soft_furnishings", "家电": "appliances", "智能": "smart_home",
        "其他": "others",
    }
    total = 0
    for r in range(7, 23):
        name = ws5.cell(row=r, column=3).value
        price_val = ws5.cell(row=r, column=4).value
        if not name:
            continue
        price = 0
        if isinstance(price_val, (int, float)):
            price = float(price_val)
        elif price_val and isinstance(price_val, str):
            try:
                price = float(price_val.replace(",", ""))
            except ValueError:
                pass
        key = None
        for cn, en in budget_map.items():
            if cn in name:
                key = en
                break
        if key:
            budget[key] = {"item": name, "amount": price}
            total += price
    budget["total"] = round(total, 2)
    result["budget"] = budget

    # ---- Sheet 6 ----
    ws6 = wb["特殊需求"]
    special = OrderedDict()
    key_map_6 = {
        "智能家居": "smart_home", "无障碍": "accessibility", "儿童安全": "child_safety",
        "环保": "eco_standard", "风水": "feng_shui", "绿植": "plants",
        "特殊材质": "special_materials", "施工时间": "construction_timing",
        "其他补充": "other_notes",
    }
    for r in range(6, 15):
        label = ws6.cell(row=r, column=2).value
        value = ws6.cell(row=r, column=3).value
        if not label:
            continue
        val = value if value and not str(value).startswith("（") else ""
        for cn, en in key_map_6.items():
            if cn in str(label):
                special[en] = val
                break
        else:
            special[str(label)] = val
    result["special_requirements"] = special

    return result


def parse_survey_to_json(filepath, output_path=None):
    """解析 Excel 并保存 JSON"""
    data = parse_survey(filepath)
    if output_path:
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    return data


if __name__ == "__main__":
    import sys
    inp = sys.argv[1] if len(sys.argv) > 1 else "templates/设计需求问卷模板.xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else "templates/设计条件.json"
    result = parse_survey_to_json(inp, out)
    print(json.dumps(result, ensure_ascii=False, indent=2))
